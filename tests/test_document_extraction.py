import io
import os
import tempfile
import unittest
from unittest.mock import patch

import app as feedback_app
import document_extraction as doc_ext


def _make_minimal_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Project Alpha Report")
    doc.add_paragraph(
        "On 1 January 2024 the project was approved by the board. "
        "A key milestone was reached on 15 June 2024 at 10:00."
    )
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def _make_minimal_pdf() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    buf.seek(0)
    return buf.read()


class ExtractTextTests(unittest.TestCase):
    def test_extract_text_from_docx(self):
        file_bytes = _make_minimal_docx()
        text = doc_ext.extract_text(file_bytes, "report.docx")
        self.assertIn("Project Alpha Report", text)

    def test_extract_text_from_pdf(self):
        file_bytes = _make_minimal_pdf()
        text = doc_ext.extract_text(file_bytes, "report.pdf")
        self.assertIsInstance(text, str)

    def test_extract_text_rejects_unsupported_format(self):
        with self.assertRaises(ValueError) as ctx:
            doc_ext.extract_text(b"data", "report.txt")
        self.assertIn("Unsupported", str(ctx.exception))


class NormaliseKeyEventsTests(unittest.TestCase):
    def test_normalises_valid_events(self):
        raw = [
            {"event": "Kick-off", "date": "2024-01-01", "time": "09:00"},
            {"event": "Launch", "date": None, "time": None},
        ]
        result = doc_ext._normalise_key_events(raw)
        self.assertEqual(len(result), 2)
        self.assertEqual(result[0]["event"], "Kick-off")
        self.assertEqual(result[0]["date"], "2024-01-01")
        self.assertEqual(result[1]["date"], None)

    def test_skips_items_without_event_text(self):
        raw = [{"event": "", "date": "2024-01-01", "time": None}]
        result = doc_ext._normalise_key_events(raw)
        self.assertEqual(result, [])

    def test_returns_empty_list_for_non_list(self):
        self.assertEqual(doc_ext._normalise_key_events(None), [])
        self.assertEqual(doc_ext._normalise_key_events("not a list"), [])


class MockDocumentFormatterTests(unittest.TestCase):
    def test_format_returns_expected_keys(self):
        formatter = doc_ext.MockDocumentFormatter()
        result = formatter.format("Annual review of the infrastructure programme")
        self.assertIn("title", result)
        self.assertIn("summary", result)
        self.assertIn("key_events", result)
        self.assertIsInstance(result["key_events"], list)


class TemplateDocumentRendererTests(unittest.TestCase):
    def _create_template(self, content: str) -> str:
        from docx import Document

        fd, path = tempfile.mkstemp(suffix=".docx")
        os.close(fd)
        doc = Document()
        doc.add_paragraph(content)
        doc.save(path)
        return path

    def test_render_replaces_placeholders(self):
        from docx import Document

        path = self._create_template("{{TITLE}}\n{{SUMMARY}}\n{{KEY_EVENTS}}")
        renderer = doc_ext.TemplateDocumentRenderer(template_path=path)
        payload = {
            "title": "Alpha Report",
            "summary": "Summary text.",
            "key_events": [{"event": "Kickoff", "date": "2024-01-01", "time": None}],
        }
        output = renderer.render(payload)
        doc = Document(io.BytesIO(output))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Alpha Report", text)
        self.assertIn("Summary text.", text)
        self.assertIn("Kickoff", text)

    def test_render_fails_when_template_missing_placeholders(self):
        path = self._create_template("Only title {{TITLE}}")
        renderer = doc_ext.TemplateDocumentRenderer(template_path=path)
        payload = {
            "title": "Alpha Report",
            "summary": "Summary text.",
            "key_events": [{"event": "Kickoff", "date": "2024-01-01", "time": None}],
        }
        with self.assertRaises(ValueError) as ctx:
            renderer.render(payload)
        self.assertIn("Missing placeholders", str(ctx.exception))

    def test_render_fails_for_missing_required_payload_fields(self):
        path = self._create_template("{{TITLE}} {{SUMMARY}} {{KEY_EVENTS}}")
        renderer = doc_ext.TemplateDocumentRenderer(template_path=path)
        with self.assertRaises(ValueError) as ctx:
            renderer.render({"title": "", "summary": "Summary text.", "key_events": []})
        self.assertIn("non-empty document title", str(ctx.exception))


class DocumentExtractionRouteTests(unittest.TestCase):
    def setUp(self):
        feedback_app.app.config["TESTING"] = True
        self.client = feedback_app.app.test_client()
        self._env = os.environ.copy()
        os.environ["DEMO_USE_MOCK_ANALYZERS"] = "true"

    def tearDown(self):
        os.environ.clear()
        os.environ.update(self._env)

    def test_get_renders_page_with_output_template(self):
        response = self.client.get("/document-extraction")
        self.assertEqual(response.status_code, 200)
        body = response.get_data(as_text=True)
        self.assertIn("Document Extraction and Formatting", body)
        self.assertIn("Document Title", body)
        self.assertIn("Summary", body)
        self.assertIn("Key Events", body)

    def test_sidebar_nav_present_on_feedback_page(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        body = response.get_data(as_text=True)
        self.assertIn("Feedback Analysis", body)
        self.assertIn("Document Extraction and Formatting", body)
        self.assertIn("/document-extraction", body)

    def test_post_docx_returns_mock_result(self):
        file_bytes = _make_minimal_docx()
        response = self.client.post(
            "/document-extraction",
            data={"document_file": (io.BytesIO(file_bytes), "report.docx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        body = response.get_data(as_text=True)
        self.assertIn("Extracted information", body)
        self.assertIn("Document Title", body)
        self.assertIn("Download reformatted Word document", body)
        self.assertIn("Template formatting complete.", body)

    def test_post_pdf_returns_mock_result(self):
        file_bytes = _make_minimal_pdf()
        with patch.object(
            doc_ext, "extract_text_from_pdf", return_value="Project Alpha PDF content for testing"
        ):
            response = self.client.post(
                "/document-extraction",
                data={"document_file": (io.BytesIO(file_bytes), "report.pdf")},
                content_type="multipart/form-data",
            )
        self.assertEqual(response.status_code, 200)
        body = response.get_data(as_text=True)
        self.assertIn("Extracted information", body)
        self.assertIn("Download reformatted Word document", body)

    def test_post_no_file_returns_error(self):
        response = self.client.post(
            "/document-extraction",
            data={},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        body = response.get_data(as_text=True)
        self.assertIn("Please upload", body)

    def test_post_unsupported_format_returns_error(self):
        response = self.client.post(
            "/document-extraction",
            data={"document_file": (io.BytesIO(b"data"), "notes.txt")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        body = response.get_data(as_text=True)
        self.assertIn("Unsupported file format", body)

    def test_post_uses_azure_openai_when_mock_disabled(self):
        os.environ["DEMO_USE_MOCK_ANALYZERS"] = "false"
        os.environ["AZURE_OPENAI_ENDPOINT"] = "https://fake.openai.azure.com"
        os.environ["LANGUAGE_MODEL_DEPLOYMENT_NAME"] = "gpt-4-1-mini"

        mock_result = {
            "title": "Alpha Report",
            "summary": "A short summary.",
            "key_events": [{"event": "Kickoff", "date": "2024-01-01", "time": None}],
        }
        file_bytes = _make_minimal_docx()

        with patch.object(doc_ext.DocumentFormatter, "format", return_value=mock_result):
            response = self.client.post(
                "/document-extraction",
                data={"document_file": (io.BytesIO(file_bytes), "report.docx")},
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200)
        body = response.get_data(as_text=True)
        self.assertIn("Alpha Report", body)
        self.assertIn("A short summary.", body)
        self.assertIn("Kickoff", body)
        self.assertIn("Download reformatted Word document", body)

    def test_post_shows_template_failure_status_when_template_missing(self):
        os.environ["DOCUMENT_TEMPLATE_PATH"] = "/tmp/does-not-exist-template.docx"
        file_bytes = _make_minimal_docx()
        response = self.client.post(
            "/document-extraction",
            data={"document_file": (io.BytesIO(file_bytes), "report.docx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        body = response.get_data(as_text=True)
        self.assertIn("Structured extraction completed, but template formatting failed.", body)
        self.assertIn("Word template file is missing.", body)
        self.assertIn("Extracted information", body)


if __name__ == "__main__":
    unittest.main()
