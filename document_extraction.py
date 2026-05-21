import io
import logging
import os
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional

from feedback_analysis import _build_service_headers, _post_with_retry, _safe_parse_json

_logger = logging.getLogger(__name__)

MAX_DOCUMENT_CHARS = 50_000
DEFAULT_TEMPLATE_PATH = (
    Path(__file__).resolve().parent / "word_templates" / "document_reformat_template.docx"
)
TEMPLATE_PLACEHOLDERS = {
    "title": "{{TITLE}}",
    "summary": "{{SUMMARY}}",
    "key_events": "{{KEY_EVENTS}}",
}

_SYSTEM_PROMPT = (
    "You are a document analyst. Extract structured information from the provided document text. "
    "Return JSON with exactly these keys:\n"
    "- title: the document title (read from the document or generated from the content, as a short string)\n"
    "- summary: a concise summary of the document in exactly 1-2 sentences\n"
    "- key_events: an array of objects in chronological order, each with keys: "
    "event (brief description of the event), date (date string or null if not provided), "
    "time (time string or null if not provided). "
    "Include only major events. If no events are present, return an empty array."
)


def extract_text_from_docx(file_bytes: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(file_bytes))
    return "\n".join(para.text for para in doc.paragraphs if para.text.strip())


def extract_text_from_pdf(file_bytes: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text.strip())
    return "\n".join(pages)


def extract_text(file_bytes: bytes, filename: str) -> str:
    filename_lower = filename.lower()
    if filename_lower.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    if filename_lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    raise ValueError("Unsupported file format. Please upload a .docx or .pdf file.")


def _normalise_key_events(raw: Any) -> List[Dict[str, Optional[str]]]:
    if not isinstance(raw, list):
        return []
    result = []
    for item in raw:
        if not isinstance(item, dict):
            continue
        event_text = str(item.get("event", "")).strip()
        if not event_text:
            continue
        result.append(
            {
                "event": event_text,
                "date": str(item["date"]).strip() if item.get("date") else None,
                "time": str(item["time"]).strip() if item.get("time") else None,
            }
        )
    return result


@dataclass
class DocumentFormatter:
    endpoint: str
    deployment: str
    api_version: str = "2025-01-01-preview"
    api_key: str = ""

    def format(self, text: str) -> Dict[str, Any]:
        if not self.endpoint or not self.deployment:
            raise ValueError(
                "Azure OpenAI configuration is missing. "
                "Set AZURE_OPENAI_ENDPOINT and a deployment name."
            )

        truncated = text[:MAX_DOCUMENT_CHARS]
        headers = _build_service_headers(api_key=self.api_key, api_key_header="api-key")
        response = _post_with_retry(
            f"{self.endpoint.rstrip('/')}/openai/deployments/{self.deployment}"
            f"/chat/completions?api-version={self.api_version}",
            headers=headers,
            json={
                "temperature": 0,
                "messages": [
                    {"role": "system", "content": _SYSTEM_PROMPT},
                    {"role": "user", "content": truncated},
                ],
                "response_format": {"type": "json_object"},
            },
            timeout=60,
        )
        choice = response.json().get("choices", [{}])[0]
        content = choice.get("message", {}).get("content", "{}")
        parsed = _safe_parse_json(content)
        return {
            "title": str(parsed.get("title", "")).strip() or "Untitled Document",
            "summary": str(parsed.get("summary", "")).strip() or "No summary available.",
            "key_events": _normalise_key_events(parsed.get("key_events", [])),
        }


@dataclass
class MockDocumentFormatter:
    def format(self, text: str) -> Dict[str, Any]:
        words = text.split()
        title = " ".join(words[:6]).rstrip(".,;:") if words else "Sample Document"
        return {
            "title": title,
            "summary": (
                "This is a demo summary of the uploaded document. "
                "The document has been processed by the mock analyzer."
            ),
            "key_events": [
                {"event": "Document created", "date": "2024-01-01", "time": None},
                {"event": "Key milestone reached", "date": "2024-06-15", "time": "10:00"},
            ],
        }


def resolve_template_path(template_path: str = "") -> Path:
    candidate = template_path.strip()
    if not candidate:
        candidate = os.getenv("DOCUMENT_TEMPLATE_PATH", "").strip()
    if candidate:
        return Path(candidate)
    return DEFAULT_TEMPLATE_PATH


def _format_key_events_text(key_events: List[Dict[str, Optional[str]]]) -> str:
    if not key_events:
        return "No key events identified."
    lines = []
    for event in key_events:
        event_name = str(event.get("event", "")).strip()
        if not event_name:
            continue
        date = str(event.get("date", "")).strip() if event.get("date") else ""
        time = str(event.get("time", "")).strip() if event.get("time") else ""
        if date and time:
            lines.append(f"- {event_name} ({date} at {time})")
        elif date:
            lines.append(f"- {event_name} ({date})")
        elif time:
            lines.append(f"- {event_name} ({time})")
        else:
            lines.append(f"- {event_name}")
    return "\n".join(lines) if lines else "No key events identified."


def _get_template_text_nodes(doc: Any) -> List[Any]:
    nodes = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                nodes.extend(cell.paragraphs)
    return nodes


def _replace_placeholder_in_doc(doc: Any, placeholder: str, replacement: str) -> int:
    replaced = 0
    for paragraph in _get_template_text_nodes(doc):
        replaced_in_runs = False
        for run in paragraph.runs:
            if placeholder in run.text:
                run.text = run.text.replace(placeholder, replacement)
                replaced_in_runs = True
        if replaced_in_runs:
            replaced += 1
            continue
        if placeholder in paragraph.text:
            paragraph.text = paragraph.text.replace(placeholder, replacement)
            replaced += 1
    return replaced


def _template_contains_placeholder(doc: Any, placeholder: str) -> bool:
    return any(placeholder in node.text for node in _get_template_text_nodes(doc))


def validate_template_payload(payload: Dict[str, Any]) -> Dict[str, Any]:
    title = str(payload.get("title", "")).strip()
    summary = str(payload.get("summary", "")).strip()
    key_events = _normalise_key_events(payload.get("key_events", []))

    if not title:
        raise ValueError("Template formatting requires a non-empty document title.")
    if not summary:
        raise ValueError("Template formatting requires a non-empty summary.")

    return {"title": title, "summary": summary, "key_events": key_events}


@dataclass
class TemplateDocumentRenderer:
    template_path: str = ""

    def render(self, payload: Dict[str, Any]) -> bytes:
        from docx import Document

        validated = validate_template_payload(payload)
        resolved_template_path = resolve_template_path(self.template_path)

        if not resolved_template_path.exists():
            raise ValueError(
                "Word template file is missing. "
                "Set the DOCUMENT_TEMPLATE_PATH environment variable "
                f"or add template at {resolved_template_path}."
            )
        if resolved_template_path.suffix.lower() != ".docx":
            raise ValueError("Word template must be a .docx file.")

        try:
            doc = Document(str(resolved_template_path))
        except Exception as exc:
            raise ValueError(
                "Word template could not be opened. "
                f"Use a valid .docx template file at {resolved_template_path}."
            ) from exc

        missing_placeholders = [
            placeholder
            for placeholder in TEMPLATE_PLACEHOLDERS.values()
            if not _template_contains_placeholder(doc, placeholder)
        ]
        if missing_placeholders:
            raise ValueError(
                "Word template is malformed. Missing placeholders: "
                + ", ".join(missing_placeholders)
            )

        replacements = {
            TEMPLATE_PLACEHOLDERS["title"]: validated["title"],
            TEMPLATE_PLACEHOLDERS["summary"]: validated["summary"],
            TEMPLATE_PLACEHOLDERS["key_events"]: _format_key_events_text(validated["key_events"]),
        }

        for placeholder, replacement in replacements.items():
            _replace_placeholder_in_doc(doc, placeholder, replacement)

        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        return output.read()
